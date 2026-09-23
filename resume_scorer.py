"""
Rule-based resume scorer (no LLM, same resume -> same score every time).

Use from the terminal:
    python resume_scorer.py resume.pdf

Needs:  pip install pypdf python-docx

Score out of 100:
    Contact & links            10
    Sections present           25
    Skills                     15
    Project/experience bullets 25
    Length & format            15
    Education details          10

PRIVACY: the result contains NO email, phone number or raw resume text.
Only yes/no flags, counts, skill names and (at most 3) short weak-bullet examples.
"""

import math
import re
import sys
from pathlib import Path

MAX_FILE_BYTES = 5 * 1024 * 1024
MAX_PDF_PAGES = 10

# ---------------------------------------------------------------------------
# Word lists
# ---------------------------------------------------------------------------
SECTION_HEADINGS = {
    "education": ["education", "academic", "academics", "educational qualification",
                  "educational qualifications", "academic qualification",
                  "academic qualifications", "qualification", "qualifications",
                  "academic background"],
    "skills": ["skills", "technical skills", "key skills", "core competencies",
               "technologies", "tech stack", "skills summary", "skill set",
               "technical proficiency", "technical expertise"],
    "projects": ["projects", "academic projects", "personal projects", "project work",
                 "projects undertaken", "key projects", "major projects", "project"],
    "experience": ["experience", "work experience", "internship", "internships",
                   "internship experience", "professional experience", "training",
                   "industrial training", "work history", "employment",
                   "experience and internships", "internships and training"],
    "certifications": ["certifications", "certification", "certificates", "certificate",
                       "courses", "licenses and certifications", "online courses",
                       "courses and certifications", "training and certifications"],
    "achievements": ["achievements", "awards", "honors", "honours", "accomplishments",
                     "extracurricular", "extracurricular activities",
                     "positions of responsibility", "awards and achievements",
                     "achievements and awards", "leadership"],
    "summary": ["summary", "objective", "career objective", "profile",
                "professional summary", "about me", "career summary"],
}

SKILL_KEYWORDS = [
    # languages
    "python", "java", "javascript", "typescript", "c++", "c#", "kotlin", "swift",
    "php", "matlab", "dart", "golang", "rust", "sql", "html", "css",
    # web
    "react", "angular", "vue", "node.js", "nodejs", "express", "django", "flask",
    "fastapi", "spring boot", "spring", ".net", "bootstrap", "tailwind", "next.js",
    "rest api", "restful", "graphql", "microservices",
    # data / db
    "mysql", "postgresql", "mongodb", "oracle", "sqlite", "redis", "firebase",
    "pandas", "numpy", "matplotlib", "power bi", "tableau", "excel", "dbms",
    # ai / ml
    "machine learning", "deep learning", "tensorflow", "pytorch", "scikit-learn",
    "nlp", "opencv", "langchain", "generative ai", "llm", "rag", "data science",
    # cloud / devops / tools
    "azure", "aws", "gcp", "docker", "kubernetes", "git", "github", "linux",
    "jenkins", "ci/cd", "terraform", "postman", "jira", "figma", "vs code",
    # cs fundamentals
    "data structures", "algorithms", "dsa", "oop", "oops", "operating systems",
    "computer networks", "system design", "design patterns",
    # testing
    "junit", "selenium", "unit testing",
]
# One-letter / very short skills are counted only as standalone items in the Skills section.
SHORT_SKILLS = {"c", "r", "go"}

ACTION_VERBS = {
    "developed", "built", "designed", "implemented", "created", "led", "managed",
    "optimized", "optimised", "reduced", "improved", "deployed", "integrated",
    "automated", "analyzed", "analysed", "engineered", "launched", "collaborated",
    "architected", "tested", "debugged", "migrated", "configured", "trained",
    "achieved", "delivered", "streamlined", "spearheaded", "coordinated",
    "organized", "organised", "presented", "secured", "enhanced", "researched",
    "established", "maintained", "resolved", "mentored", "published", "won",
    "constructed", "programmed", "wrote", "produced", "increased", "decreased",
    "boosted", "cut", "saved", "scaled", "refactored", "documented", "conducted",
    "initiated", "devised", "formulated", "executed", "simulated", "modeled",
    "modelled", "extracted", "processed", "visualized", "visualised", "applied",
    "utilized", "utilised", "leveraged", "handled", "supported", "contributed",
    "revamped", "redesigned", "customized", "customised", "compared", "evaluated",
}

DEGREE_PATTERN = re.compile(
    r"\b(b\.?\s?tech|b\.?\s?e\b|bachelor|b\.?\s?sc|bca|mca|m\.?\s?tech|m\.?\s?sc|"
    r"diploma|mba|b\.?\s?com|engineering)\b", re.I)
GPA_PATTERN = re.compile(
    r"(cgpa|gpa|percentage|aggregate|\b\d{1,2}(?:\.\d{1,2})?\s*/\s*10\b|\b\d{2}(?:\.\d+)?\s*%)", re.I)
YEAR_PATTERN = re.compile(r"\b(19|20)\d{2}\b")

EMAIL_RE = re.compile(r"[\w.+-]+@[\w-]+\.[\w.-]+")
PHONE_RE = re.compile(r"(?:\+?\d[\s().-]?){10,14}")
NUMBER_RE = re.compile(r"\d[\d,]*(?:\.\d+)?")
YEAR_ONLY_RE = re.compile(r"^(19|20)\d{2}$")

INJECTION_RE = re.compile(
    r"(ignore|disregard|forget)\s+(all\s+|any\s+|the\s+)?(previous|prior|above|earlier)"
    r"|system prompt|as an ai\b|you are (now )?(an?|the)\s"
    r"|give (me )?(a |an )?(score|rating)|score\s*(of|=|:)?\s*100",
    re.I)

BULLET_CHARS = "•●▪◦·–—-*►➤✓✔\u2022\uf0b7\uf0a7\uf0d8"


# ---------------------------------------------------------------------------
# Text extraction
# ---------------------------------------------------------------------------
def extract_text(path):
    """Return (text, pages, kind). pages is None for DOCX (unknown)."""
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(str(path))
        pages = len(reader.pages)
        parts = [(page.extract_text() or "") for page in reader.pages[:MAX_PDF_PAGES]]
        return "\n".join(parts), pages, "pdf"
    if suffix == ".docx":
        import docx
        document = docx.Document(str(path))
        lines = [p.text for p in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    lines.extend(cell.text.splitlines())
        return "\n".join(lines), None, "docx"
    raise ValueError("unsupported")


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------
def _norm(text):
    return re.sub(r"[^a-z ]", "", text.lower()).strip()


def _heading_of(line):
    """Return (section_name, rest_of_line) if the line is a section heading."""
    raw = line.strip()
    if not raw or len(raw.split()) > 7:
        return None
    head, sep, rest = raw.partition(":")
    candidate = _norm(head)
    if not candidate:
        return None
    for name, words in SECTION_HEADINGS.items():
        if candidate in words:
            # "Skills: Python, Java" -> heading with inline content
            return name, rest.strip() if sep else ""
    return None


def _split_sections(lines):
    sections, current = {}, None
    for line in lines:
        found = _heading_of(line)
        if found:
            current = found[0]
            sections.setdefault(current, [])
            if found[1]:
                sections[current].append(found[1])
        elif current:
            sections[current].append(line)
    return sections


def _strip_bullet(line):
    return line.strip().lstrip(BULLET_CHARS + " \t").strip()


def _first_word(text):
    match = re.match(r"[A-Za-z]+", text)
    return match.group(0).lower() if match else ""


def _has_number(text):
    for token in NUMBER_RE.findall(text):
        clean = token.replace(",", "")
        if not YEAR_ONLY_RE.match(clean):
            return True
    return "%" in text


def _find_skills(text, skills_section_lines):
    lowered = text.lower()
    found = set()
    for skill in SKILL_KEYWORDS:
        pattern = r"(?<![a-z0-9+#.])" + re.escape(skill) + r"(?![a-z0-9+#])"
        if re.search(pattern, lowered):
            found.add(skill)
    tokens = set()
    for line in skills_section_lines:
        for part in re.split(r"[,|/•;:()\u2022\t]+|\s{2,}", line.lower()):
            tokens.add(part.strip())
    for short in SHORT_SKILLS:
        if short in tokens:
            found.add(short)
    return sorted(found)


def _grade(score):
    if score >= 88:
        return "Excellent"
    if score >= 75:
        return "Good"
    if score >= 55:
        return "Needs work"
    return "Weak"


def _cat(name, score, max_score, notes):
    return {"name": name, "score": round(score, 1), "max": max_score, "notes": notes}


# ---------------------------------------------------------------------------
# Main analysis
# ---------------------------------------------------------------------------
def analyze_text(text, pages=None):
    """Score resume text. pages=None -> estimated from word count."""
    raw_lines = [ln.strip() for ln in text.replace("\r", "\n").split("\n")]
    raw_lines = [ln for ln in raw_lines if ln]

    # Prompt-injection guard: lines that try to instruct the scorer are ignored.
    suspicious = [ln for ln in raw_lines if INJECTION_RE.search(ln)]
    lines = [ln for ln in raw_lines if not INJECTION_RE.search(ln)]
    clean_text = "\n".join(lines)
    words = len(clean_text.split())

    pages_estimated = pages is None
    if pages is None:
        pages = max(1, math.ceil(words / 550))

    sections = _split_sections(lines)
    categories, improvements = [], []

    def lose(points, message):
        if points > 0.05:
            improvements.append((points, message))

    # 1) Contact & links (10)
    has_email = bool(EMAIL_RE.search(clean_text))
    has_phone = bool(PHONE_RE.search(clean_text))
    lowered = clean_text.lower()
    has_linkedin = "linkedin.com" in lowered or "linkedin" in lowered
    has_github = "github.com" in lowered
    pts = 4 * has_email + 3 * has_phone + 2 * has_linkedin + 1 * has_github
    notes = []
    if not has_email:
        notes.append("Email not found"); lose(4, "Add a professional email address at the top.")
    if not has_phone:
        notes.append("Phone number not found"); lose(3, "Add your phone number at the top.")
    if not has_linkedin:
        notes.append("LinkedIn link not found"); lose(2, "Add your LinkedIn profile link.")
    if not has_github:
        notes.append("GitHub link not found"); lose(1, "Add your GitHub profile link (important for tech roles).")
    categories.append(_cat("Contact & links", pts, 10, notes))

    # 2) Sections (25)
    weights = {"education": 6, "skills": 6, "projects": 7, "experience": 4, "certifications": 2}
    tips = {
        "education": "Add an Education section (degree, college, CGPA, year).",
        "skills": "Add a Skills section listing languages, frameworks, databases and tools.",
        "projects": "Add a Projects section with 2-3 projects (what you built, tech used, result).",
        "experience": "Add Internships/Experience/Training. If you have none, add freelance, open-source or hackathon work.",
        "certifications": "Add relevant certifications or courses (for example Microsoft Learn, NPTEL).",
    }
    pts, notes, present = 0, [], []
    for name, weight in weights.items():
        if name in sections and sections[name]:
            pts += weight
            present.append(name)
        else:
            notes.append(f"{name.capitalize()} section missing")
            lose(weight, tips[name])
    categories.append(_cat("Sections present", pts, 25, notes))

    # 3) Skills (15)
    text_for_skills = re.sub(r"\S+@\S+|(?:https?://|www\.)\S+|\b\S+\.(?:com|in|io|dev|org)\S*", " ", clean_text)
    skills_found = _find_skills(text_for_skills, sections.get("skills", []))
    pts = min(15, len(skills_found) * 15 / 12)
    notes = [f"{len(skills_found)} recognised skills found"]
    if len(skills_found) < 12:
        lose(15 - pts, "List more relevant skills (aim for 12+): languages, frameworks, databases, cloud and tools.")
    categories.append(_cat("Skills", pts, 15, notes))

    # 4) Project / experience bullets (25)
    source_lines = sections.get("projects", []) + sections.get("experience", [])
    candidates = []
    for line in source_lines:
        stripped = _strip_bullet(line)
        if len(stripped.split()) >= 5 and not _heading_of(stripped):
            candidates.append(stripped)
    action, quantified, weak_examples = 0, 0, []
    for bullet in candidates:
        has_verb = _first_word(bullet) in ACTION_VERBS
        has_num = _has_number(bullet)
        action += has_verb
        quantified += has_num
        if not has_verb and not has_num and len(weak_examples) < 3:
            weak_examples.append(bullet[:110])
    total = len(candidates)
    if total:
        verb_pts = 10 * min(1.0, (action / total) / 0.7)
        num_pts = 10 * min(1.0, (quantified / total) / 0.4)
        vol_pts = 5 * min(1.0, total / 6)
    else:
        verb_pts = num_pts = vol_pts = 0
    notes = [f"{total} bullet lines analysed", f"{action} start with an action verb",
             f"{quantified} contain numbers/results"]
    if total == 0:
        lose(25, "Add bullet points under Projects/Experience describing what you did and the result.")
    else:
        lose(10 - verb_pts, "Start bullets with strong action verbs (Built, Designed, Implemented, Reduced...).")
        lose(10 - num_pts, "Add measurable results to bullets (for example: reduced load time by 30%, served 500+ users).")
        lose(5 - vol_pts, "Add more detail: aim for 6+ bullets across projects/experience.")
    categories.append(_cat("Project/experience bullets", verb_pts + num_pts + vol_pts, 25, notes))

    # 5) Length & format (15)
    if pages == 1:
        page_pts = 8
    elif pages == 2:
        page_pts = 4
        lose(4, "Keep the resume to 1 page (freshers) - shorten or cut older/less relevant points.")
    else:
        page_pts = 1
        lose(7, "Resume is too long. Reduce it to 1 page.")
    if 250 <= words <= 700:
        word_pts = 7
    elif 150 <= words < 250 or 700 < words <= 900:
        word_pts = 4
        lose(3, "Adjust length: aim for about 250-700 words.")
    else:
        word_pts = 1
        lose(6, "Adjust length: aim for about 250-700 words.")
    notes = [f"{pages} page(s){' (estimated)' if pages_estimated else ''}", f"{words} words"]
    categories.append(_cat("Length & format", page_pts + word_pts, 15, notes))

    # 6) Education details (10)
    edu_lines = sections.get("education") or lines
    edu_text = "\n".join(edu_lines)
    has_degree = bool(DEGREE_PATTERN.search(edu_text))
    has_gpa = bool(GPA_PATTERN.search(edu_text))
    has_year = bool(YEAR_PATTERN.search(edu_text))
    pts = 4 * has_degree + 4 * has_gpa + 2 * has_year
    notes = []
    if not has_degree:
        notes.append("Degree name not found"); lose(4, "Mention your degree clearly (for example B.Tech in Computer Science).")
    if not has_gpa:
        notes.append("CGPA/percentage not found"); lose(4, "Mention your CGPA or percentage in the Education section.")
    if not has_year:
        notes.append("Year not found"); lose(2, "Mention start/end or expected graduation year.")
    categories.append(_cat("Education details", pts, 10, notes))

    total_score = round(sum(c["score"] for c in categories))
    improvements.sort(key=lambda item: -item[0])
    top_improvements = [message for _, message in improvements][:6]
    strengths = [f"{c['name']}: {c['score']}/{c['max']}" for c in categories if c["score"] >= 0.8 * c["max"]]

    notes_out = ["Score is rule-based (same resume gives the same score); it is not an ATS or recruiter guarantee."]
    if suspicious:
        notes_out.append("Some lines looked like instructions to an AI and were ignored in scoring.")

    return {
        "status": "OK",
        "total_score": total_score,
        "grade": _grade(total_score),
        "categories": categories,
        "strengths": strengths,
        "improvements": top_improvements,
        "found": {
            "sections": present,
            "missing_sections": [s for s in weights if s not in present],
            "skills": skills_found,
            "pages": pages,
            "pages_estimated": pages_estimated,
            "word_count": words,
            "bullets_analyzed": total,
            "action_verb_bullets": action,
            "quantified_bullets": quantified,
        },
        "weak_bullet_examples": weak_examples,
        "suspicious_text_found": bool(suspicious),
        "notes": notes_out,
    }


def score_resume_file(path):
    """Score a resume file. Never raises: returns a dict with a status."""
    try:
        path = Path(path)
        if not path.is_file():
            return {"status": "FILE_NOT_FOUND",
                    "message": "No resume file found. Ask the student to upload a PDF or DOCX resume."}
        if path.stat().st_size > MAX_FILE_BYTES:
            return {"status": "FILE_TOO_LARGE", "message": "Resume file is larger than 5 MB."}
        if path.suffix.lower() not in {".pdf", ".docx"}:
            return {"status": "UNSUPPORTED_FILE", "message": "Only PDF and DOCX resumes are supported."}
        text, pages, _ = extract_text(path)
        if len(text.split()) < 50:
            return {"status": "NO_TEXT_FOUND",
                    "message": "Very little text could be read. The PDF may be a scanned image; "
                               "use a text-based PDF or DOCX."}
        return analyze_text(text, pages)
    except Exception as exc:  # keep the agent alive; it will explain the failure
        return {"status": "ERROR", "message": f"Could not read the resume ({type(exc).__name__})."}


# ---------------------------------------------------------------------------
# Terminal report
# ---------------------------------------------------------------------------
def print_report(result, show_bullets=False):
    if result.get("status") != "OK":
        print("Status:", result.get("status"), "-", result.get("message", ""))
        return
    print(f"\nRESUME SCORE: {result['total_score']}/100  ({result['grade']})\n")
    for c in result["categories"]:
        print(f"  {c['name']:<28} {c['score']:>5}/{c['max']}")
        for note in c["notes"]:
            print(f"       - {note}")
    print("\nTop improvements:")
    for i, item in enumerate(result["improvements"], 1):
        print(f"  {i}. {item}")
    if show_bullets and result["weak_bullet_examples"]:
        print("\nBullets that need stronger wording/results:")
        for b in result["weak_bullet_examples"]:
            print("   -", b)
    found = result["found"]
    print("\nSections found:", ", ".join(found["sections"]) or "none")
    print("Skills found  :", ", ".join(found["skills"]) or "none")
    for note in result["notes"]:
        print("Note:", note)


if __name__ == "__main__":
    args = [a for a in sys.argv[1:] if not a.startswith("--")]
    if len(args) != 1:
        raise SystemExit("Usage: python resume_scorer.py resume.pdf [--show-bullets]")
    print_report(score_resume_file(args[0]), show_bullets="--show-bullets" in sys.argv)
